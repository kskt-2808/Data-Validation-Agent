"""Excel export: Summary, Ledger and Method sheets built from a validation result."""
from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
HEADER_FONT = Font(bold=True, color="FFFFFF")
STATUS_FILLS = {s: PatternFill("solid", fgColor=c) for s, c in
                {"PASS": "DCF3E6", "WARN": "FFF1CC", "NO_DATA": "FDE4E1", "ERROR": "EAECF0"}.items()}
STATUS_LABELS = {"NO_DATA": "NO DATA"}


def build_workbook(result: dict) -> bytes:
    wb = Workbook()
    _summary_sheet(wb.active, result)
    _ledger_sheet(wb.create_sheet("Ledger"), result)
    _method_sheet(wb.create_sheet("Method"), result)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def filename(result: dict) -> str:
    p = result["params"]
    return f"newton_{result['recipe']['id']}_{p['from']}_to_{p['to']}.xlsx"


def _summary_sheet(ws, result: dict) -> None:
    recipe, p = result["recipe"], result["params"]
    ws.title = "Summary"
    ws.append(["Newton — Data Validation Report"])
    ws["A1"].font = Font(bold=True, size=14)
    for label, value in [
        ("Calculation", recipe["label"]),
        ("Formula", recipe["formula"]),
        ("Date range", f"{p['from']} to {p['to']}"),
        ("Shift window", f"{p['shift']} ({result['timezone']})"),
        ("Generated", result["generatedAt"]),
    ]:
        ws.append([label, value])
    for group in result["summary"]:
        unit = group["unit"]
        ws.append([])
        ws.append([f"Figures in {unit}" if unit else "Figures"])
        ws.cell(ws.max_row, 1).font = Font(bold=True)
        if recipe["aggregate"] == "sum":
            ws.append([recipe["totalLabel"], group["total"], unit])
        ws.append([recipe["avgLabel"], group["mean"], unit])
        peak = group["peak"]
        ws.append([recipe["peakLabel"], peak and peak["value"], unit,
                   peak and f"{peak['date']} · {peak['devID']} / {peak['sensor']}"])
        c = group["counts"]
        ws.append(["Valid shifts", c["PASS"] + c["WARN"], "", f"{c['WARN']} with warnings"])
        ws.append(["No-data shifts", c["NO_DATA"]])
        ws.append(["Fetch errors", c["ERROR"]])
    for row in ws.iter_rows(min_col=2, max_col=2):
        for cell in row:
            if isinstance(cell.value, float):
                cell.number_format = "#,##0.00"
    _widths(ws, [28, 48, 10, 36])


def _ledger_columns(recipe: dict) -> list[tuple[str, str, str]]:
    """(header, row key, kind) with value/time pairs split and audit columns appended."""
    cols = []
    for c in recipe["columns"]:
        if c["type"] == "status":
            continue
        cols.append((c["label"], c["key"], c["type"]))
        if c.get("timeKey"):
            cols.append((c["label"].split(" (")[0] + " Time", c["timeKey"], "time"))
        if c["type"] == "output":
            cols.append(("Unit", "unit", "text"))
    cols += [("Status", "status", "status"), ("Notes", "notes", "notes"),
             ("Device Name", "devName", "text"), ("Sensor Name", "sensorName", "text"),
             ("m", "m", "factor"), ("c", "c", "factor"), ("Factor Source", "factorSource", "text"),
             ("Window Start", "windowStart", "time"), ("Window End", "windowEnd", "time")]
    return cols


def _ledger_sheet(ws, result: dict) -> None:
    cols = _ledger_columns(result["recipe"])
    ws.append([h for h, _, _ in cols])
    for cell in ws[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    for row in result["rows"]:
        values = []
        for _, key, kind in cols:
            v = row.get(key)
            if kind == "status":
                v = STATUS_LABELS.get(v, v)
            elif kind == "notes":
                v = "; ".join(v or [])
            elif kind == "time" and v:
                v = v[:19].replace("T", " ")
            values.append(v)
        ws.append(values)
        for (_, key, kind), cell in zip(cols, ws[ws.max_row]):
            if kind in ("number", "output"):
                cell.number_format = "#,##0.00"
            elif kind == "factor":
                cell.number_format = "0.######"
            elif kind == "status":
                cell.fill = STATUS_FILLS.get(row["status"], PatternFill())
    ws.auto_filter.ref = ws.dimensions
    _widths(ws, [max(12, min(40, len(h) + 4)) for h, _, _ in cols])


def _method_sheet(ws, result: dict) -> None:
    recipe, p = result["recipe"], result["params"]
    lines = [
        ("Calculation", recipe["label"]),
        ("Formula", recipe["formula"]),
        ("Method", recipe["method"]),
        ("Data source", f"{result['source']} — raw readings; Newton applies m and c itself "
                        "(see the m, c and Factor Source columns on the Ledger sheet)."),
        ("Shift window", f"{p['shift']} in {result['timezone']}, half-open [start, end): a reading "
                         "stamped exactly at the end belongs to the next shift."),
        ("Max gap", f"{p['maxGapMinutes']:g} min. A reading holds for at most this long; longer "
                    "silences count as unknown time."),
        ("Unit scale", f"÷ {p['unitScale']}" if p["unitScale"] != 1 else "none (native unit)"),
        ("Statuses", "PASS = computed with no concerns. WARN = computed, but see Notes. "
                     "NO DATA = not enough readings to compute (never reported as zero). "
                     "ERROR = the platform did not return a usable response."),
    ]
    for label, text in lines:
        ws.append([label, text])
        ws.cell(ws.max_row, 1).font = Font(bold=True)
        ws.cell(ws.max_row, 2).alignment = Alignment(wrap_text=True, vertical="top")
    _widths(ws, [16, 110])


def _widths(ws, widths: list[int]) -> None:
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
