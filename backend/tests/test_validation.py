import json
import sys
import unittest
from io import BytesIO
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from openpyxl import load_workbook  # noqa: E402

from iosense import UpstreamError  # noqa: E402
from report import build_workbook  # noqa: E402
from validation import run_validation  # noqa: E402

FIXTURES = Path(__file__).parent / "fixtures"
SHIFT_START_MS = 1789349400000  # 2026-09-14 07:00 IST
MIN = 60_000


def _device(dev_id, *sensors):
    return {"devID": dev_id, "devName": dev_id, "devType": "Energymeter", "sensors": list(sensors)}


ENERGY = {"id": "D30", "name": "Active Energy", "unit": "kWh", "m": 1.0, "c": 0.0}
CURRENT = {"id": "D6", "name": "Average Current", "unit": "A", "m": 1.0, "c": 0.0}
DEVICES = {"EVOEM_C1": _device("EVOEM_C1", ENERGY, CURRENT), "BROKEN": _device("BROKEN", ENERGY)}


class FakeClient:
    """Serves fixture readings; devices without a fixture fail like the platform does."""

    def __init__(self):
        self.series = {}
        for sensor in ("d30", "d6"):
            fixture = json.loads((FIXTURES / f"evoem_c1_{sensor}_2026-09-14.json").read_text())
            self.series[("EVOEM_C1", fixture["sensor"])] = fixture["points"]
        self.calls = []

    def raw_series(self, dev_id, sensor, start_ms, end_ms):
        self.calls.append((dev_id, sensor, start_ms, end_ms))
        if (dev_id, sensor) not in self.series:
            raise UpstreamError("Error in fetch data")
        return [(t, v) for t, v in self.series[(dev_id, sensor)] if start_ms <= t <= end_ms]


def request(**overrides):
    body = {"formula": "consumption_delta", "from": "2026-09-13", "to": "2026-09-14",
            "shiftStart": "07:00", "shiftEnd": "07:00",
            "targets": [{"devID": "EVOEM_C1", "sensor": "D30"}, {"devID": "BROKEN", "sensor": "D30"}]}
    body.update(overrides)
    return body


class RunValidation(unittest.TestCase):
    def setUp(self):
        self.client = FakeClient()

    def test_consumption_report_end_to_end(self):
        result = run_validation(self.client, DEVICES, request())
        rows = {(r["date"], r["devID"]): r for r in result["rows"]}

        good = rows[("2026-09-14", "EVOEM_C1")]
        self.assertEqual(good["status"], "PASS")
        self.assertAlmostEqual(good["value"], 1268.0)
        self.assertEqual(good["first_time"], "2026-09-14T07:00:31+05:30")
        self.assertEqual(good["last_time"], "2026-09-15T06:58:51+05:30")
        self.assertEqual(rows[("2026-09-13", "EVOEM_C1")]["status"], "NO_DATA")
        self.assertIsNone(rows[("2026-09-13", "EVOEM_C1")]["value"])
        self.assertEqual(rows[("2026-09-14", "BROKEN")]["status"], "ERROR")

        [group] = result["summary"]
        self.assertAlmostEqual(group["total"], 1268.0)
        self.assertEqual(group["counts"], {"PASS": 1, "WARN": 0, "NO_DATA": 1, "ERROR": 2})
        self.assertIn(("EVOEM_C1", "D30", SHIFT_START_MS - 15 * MIN, SHIFT_START_MS + 24 * 60 * MIN),
                      self.client.calls)

    def test_override_and_unit_conversion_are_recorded(self):
        result = run_validation(self.client, DEVICES, request(
            to="2026-09-14", outputUnit="MWh",
            targets=[{"devID": "EVOEM_C1", "sensor": "D30", "m": 1000, "c": 0}]))
        row = next(r for r in result["rows"] if r["status"] == "PASS")
        self.assertEqual(row["factorSource"], "override")
        self.assertEqual(row["factor"], 1.0)          # m 1000, then kWh -> MWh divides by 1000
        self.assertEqual(row["unit"], "MWh")
        self.assertAlmostEqual(row["value"], 1268.0)

    def test_rejects_a_conversion_the_unit_cannot_make(self):
        with self.assertRaisesRegex(ValueError, "cannot be shown as"):
            run_validation(self.client, DEVICES, request(
                to="2026-09-14", outputUnit="mA",
                targets=[{"devID": "EVOEM_C1", "sensor": "D30"}]))

    def test_average_value_is_the_plain_mean(self):
        result = run_validation(self.client, DEVICES, request(
            formula="average_value", to="2026-09-14",
            targets=[{"devID": "EVOEM_C1", "sensor": "D6"}]))
        row = next(r for r in result["rows"] if r["date"] == "2026-09-14")
        self.assertEqual(row["unit"], "A")
        self.assertEqual(row["status"], "PASS")
        points = [v for t, v in self.client.series[("EVOEM_C1", "D6")]
                  if SHIFT_START_MS <= t < SHIFT_START_MS + 24 * 60 * MIN]
        self.assertAlmostEqual(row["value"], sum(points) / len(points))
        self.assertEqual(row["samples"], len(points))

    def test_run_hours_can_report_in_minutes(self):
        body = request(formula="run_hours", to="2026-09-14", outputUnit="min",
                       targets=[{"devID": "EVOEM_C1", "sensor": "D6", "threshold": 3}])
        minutes = next(r for r in run_validation(self.client, DEVICES, body)["rows"]
                       if r["date"] == "2026-09-14")
        body.pop("outputUnit")
        hours = next(r for r in run_validation(self.client, DEVICES, body)["rows"]
                     if r["date"] == "2026-09-14")
        self.assertEqual(minutes["unit"], "min")
        self.assertAlmostEqual(minutes["value"], hours["value"] * 60)

    def test_availability_divides_run_hours_by_planned_hours(self):
        body = request(formula="availability_ratio", to="2026-09-14", plannedHours=12,
                       targets=[{"devID": "EVOEM_C1", "sensor": "D6", "threshold": 3}])
        row = next(r for r in run_validation(self.client, DEVICES, body)["rows"]
                   if r["date"] == "2026-09-14")
        self.assertEqual(row["unit"], "%")
        self.assertEqual(row["planned_hours"], 12)
        self.assertAlmostEqual(row["value"], row["run_hours"] / 12 * 100)

    def test_load_factor_is_average_over_peak(self):
        row = next(r for r in run_validation(self.client, DEVICES, request(
            formula="load_factor", to="2026-09-14",
            targets=[{"devID": "EVOEM_C1", "sensor": "D6"}]))["rows"] if r["date"] == "2026-09-14")
        self.assertEqual(row["unit"], "%")
        self.assertAlmostEqual(row["value"], row["average"] / row["peak"] * 100)

    def test_run_hours_needs_a_threshold_for_every_device(self):
        with self.assertRaisesRegex(ValueError, "threshold"):
            run_validation(self.client, DEVICES, request(
                formula="run_hours", targets=[{"devID": "EVOEM_C1", "sensor": "D6"}]))

    def test_run_hours_row(self):
        result = run_validation(self.client, DEVICES, request(
            formula="run_hours", targets=[{"devID": "EVOEM_C1", "sensor": "D6", "threshold": 3}]))
        row = next(r for r in result["rows"] if r["date"] == "2026-09-14")
        self.assertEqual(row["unit"], "h")
        self.assertAlmostEqual(row["value"] + row["stopped_hours"] + row["unknown_hours"], 24.0)

    def test_rejects_unknown_device_and_future_formula(self):
        with self.assertRaisesRegex(ValueError, "not in your account"):
            run_validation(self.client, DEVICES, request(targets=[{"devID": "ZYDEM_D4", "sensor": "D30"}]))
        with self.assertRaisesRegex(ValueError, "unavailable"):
            run_validation(self.client, DEVICES, request(formula="specific_energy"))

    def test_excel_export(self):
        result = run_validation(self.client, DEVICES, request())
        wb = load_workbook(BytesIO(build_workbook(result)))
        self.assertEqual(wb.sheetnames, ["Summary", "Ledger", "Method"])
        ledger = wb["Ledger"]
        header = [c.value for c in ledger[1]]
        self.assertIn("First Point Time", header)
        self.assertEqual(ledger.max_row, len(result["rows"]) + 1)
        statuses = {row[header.index("Status")] for row in ledger.iter_rows(min_row=2, values_only=True)}
        self.assertEqual(statuses, {"PASS", "NO DATA", "ERROR"})


if __name__ == "__main__":
    unittest.main()


class WaterFormula(unittest.TestCase):
    """The D6 fixture stands in for a flow meter reading m3/h."""

    def setUp(self):
        self.client = FakeClient()

    def test_flow_volume_integrates_the_shift(self):
        body = request(formula="flow_volume", to="2026-09-14", rateBasisSeconds=3600,
                       outputUnitLabel="m3", targets=[{"devID": "EVOEM_C1", "sensor": "D6"}])
        row = next(r for r in run_validation(self.client, DEVICES, body)["rows"]
                   if r["date"] == "2026-09-14")
        self.assertEqual(row["status"], "PASS")
        self.assertEqual(row["unit"], "m3")                       # from the label, not the sensor
        self.assertGreater(row["value"], 0)
        self.assertAlmostEqual(row["known_hours"] + row["unknown_hours"], 24.0)

    def test_rate_basis_scales_the_total(self):
        def total(basis):
            body = request(formula="flow_volume", to="2026-09-14", rateBasisSeconds=basis,
                           targets=[{"devID": "EVOEM_C1", "sensor": "D6"}])
            return next(r for r in run_validation(self.client, DEVICES, body)["rows"]
                        if r["date"] == "2026-09-14")["value"]
        self.assertAlmostEqual(total(60), total(3600) * 60)       # per-minute vs per-hour

    def test_unit_falls_back_to_the_sensor_when_no_label(self):
        body = request(formula="flow_volume", to="2026-09-14",
                       targets=[{"devID": "EVOEM_C1", "sensor": "D6"}])
        row = next(r for r in run_validation(self.client, DEVICES, body)["rows"]
                   if r["date"] == "2026-09-14")
        self.assertEqual(row["unit"], "A")


class ParameterDefaults(unittest.TestCase):
    def setUp(self):
        self.client = FakeClient()

    def test_a_default_satisfies_a_required_parameter(self):
        body = request(formula="flow_volume", to="2026-09-14",
                       targets=[{"devID": "EVOEM_C1", "sensor": "D6"}])
        result = run_validation(self.client, DEVICES, body)          # no rateBasisSeconds sent
        self.assertEqual(result["params"]["rateBasisSeconds"], 3600)

    def test_a_required_parameter_without_a_default_is_still_enforced(self):
        with self.assertRaisesRegex(ValueError, "threshold"):
            run_validation(self.client, DEVICES, request(
                formula="run_hours", targets=[{"devID": "EVOEM_C1", "sensor": "D6"}]))
